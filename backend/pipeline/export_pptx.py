"""
export_pptx.py
--------------
Builds a PowerPoint catalog from assembled slide data.
Matches the reference template (13.33" × 7.50" widescreen, garment image layout).

Each slide dict may include a "layout" key with per-element position overrides
(already converted to inches by the export route). Missing keys fall back to defaults.
"""

import io
import os
import logging
from PIL import Image as PILImage
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

logger = logging.getLogger(__name__)

SLIDE_W_IN = 13.33
SLIDE_H_IN = 7.50


def _resolve(path: str | None) -> str | None:
    if not path:
        return None
    if path.startswith("storage/"):
        return f"/app/{path}"
    return path


def _build_specs_text(ref_number, afs, fabric, gsm) -> str:
    lines = []
    if ref_number:
        lines.append(f"REF NO   : {ref_number}")
    if afs:
        lines.append(f"AFS          : {afs}")
    if fabric:
        lines.append(f"COMP     : {fabric}")
    if gsm:
        lines.append(f"GSM        : {gsm}")
    return "\n".join(lines)


def _add_picture(slide, path, box_left_in, box_top_in, box_w_in, box_h_in, rotation=0):
    """Place image using object-fit:contain — same as the browser preview.

    Reads the image's natural dimensions, scales it to fit within the given box
    while preserving aspect ratio, then centers it in the box.
    This ensures the PPTX output matches the browser canvas exactly.
    """
    if not path or not os.path.exists(path):
        return
    try:
        with PILImage.open(path) as img:
            img_w, img_h = img.size

        img_aspect = img_w / img_h
        box_aspect = box_w_in / box_h_in

        if img_aspect > box_aspect:
            disp_w = box_w_in
            disp_h = box_w_in / img_aspect
        else:
            disp_h = box_h_in
            disp_w = box_h_in * img_aspect

        offset_x = (box_w_in - disp_w) / 2
        offset_y = (box_h_in - disp_h) / 2

        pic = slide.shapes.add_picture(
            path,
            Inches(box_left_in + offset_x),
            Inches(box_top_in + offset_y),
            Inches(disp_w),
            Inches(disp_h),
        )
        if rotation:
            pic.rotation = rotation
    except Exception as exc:
        logger.warning("export_pptx: could not add picture %s: %s", path, exc)


def _add_specs(slide, text, left_in, top_in, width_in, height_in=2.1, font_pt=10):
    if not text:
        return
    txBox = slide.shapes.add_textbox(
        Inches(left_in), Inches(top_in),
        Inches(width_in), Inches(height_in),
    )
    tf = txBox.text_frame
    tf.word_wrap = False
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.size = Pt(font_pt)
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)


def _el(lo: dict, key: str, default: dict) -> dict:
    """Return layout override for key, or default."""
    return lo.get(key) or default


LOGO_LEFT_IN  = 12.52
LOGO_TOP_IN   = 6.81
LOGO_W_IN     = 0.55
LOGO_H_IN     = 0.50


def _add_logo(slide, logo_path: str | None, pos: dict | None = None):
    if not logo_path:
        return
    resolved = _resolve(logo_path)
    if not resolved or not os.path.exists(resolved):
        return
    left = pos["left"]            if pos else LOGO_LEFT_IN
    top  = pos["top"]             if pos else LOGO_TOP_IN
    w    = pos["width"]           if pos else LOGO_W_IN
    h    = pos.get("height", LOGO_H_IN) if pos else LOGO_H_IN
    _add_picture(slide, resolved, left, top, w, h)


def _add_garment_slide(prs, slide_data: dict, logo_path: str | None = None, logo_pos: dict | None = None, font_pt: float = 10):
    layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(layout)

    front  = _resolve(slide_data.get("front_image_path"))
    back   = _resolve(slide_data.get("back_image_path"))
    detail = _resolve(slide_data.get("detail_image_path"))

    has_front  = bool(front  and os.path.exists(front))
    has_back   = bool(back   and os.path.exists(back))
    has_detail = bool(detail and os.path.exists(detail))

    specs = _build_specs_text(
        slide_data.get("ref_number"),
        slide_data.get("afs"),
        slide_data.get("fabric"),
        slide_data.get("gsm"),
    )

    lo = slide_data.get("layout") or {}

    if has_detail:
        fl = _el(lo, "front",  {"left": 0.05, "top": 0.17, "width": 4.62, "height": 6.62})
        bl = _el(lo, "back",   {"left": 4.71, "top": 0.17, "width": 4.60, "height": 6.62})
        dl = _el(lo, "detail", {"left": 9.34, "top": 0.17, "width": 3.95, "height": 2.42})
        sl = _el(lo, "specs",  {"left": 9.28, "top": 5.17, "width": 4.48, "height": 1.62})
        if has_front:
            _add_picture(slide, front,  fl["left"], fl["top"], fl["width"], fl.get("height", 6.62), fl.get("rotation", 0))
        if has_back:
            _add_picture(slide, back,   bl["left"], bl["top"], bl["width"], bl.get("height", 6.62), bl.get("rotation", 0))
        _add_picture(slide, detail, dl["left"], dl["top"], dl["width"], dl.get("height", 2.42), dl.get("rotation", 0))
        _add_specs(slide, specs, sl["left"], sl["top"], sl["width"], sl.get("height", 1.62), font_pt=font_pt)

    elif has_front and has_back:
        fl = _el(lo, "front", {"left": 0.05, "top": 0.17, "width": 4.62, "height": 6.62})
        bl = _el(lo, "back",  {"left": 4.71, "top": 0.17, "width": 4.60, "height": 6.62})
        sl = _el(lo, "specs", {"left": 9.28, "top": 5.17, "width": 4.48, "height": 1.62})
        _add_picture(slide, front, fl["left"], fl["top"], fl["width"], fl.get("height", 6.62), fl.get("rotation", 0))
        _add_picture(slide, back,  bl["left"], bl["top"], bl["width"], bl.get("height", 6.62), bl.get("rotation", 0))
        _add_specs(slide, specs, sl["left"], sl["top"], sl["width"], sl.get("height", 1.62), font_pt=font_pt)

    else:
        img_key = "front" if has_front else ("back" if has_back else "detail")
        img = front or back or detail
        il = _el(lo, img_key, {"left": 2.0, "top": 0.5, "width": 9.0, "height": 6.5})
        sl = _el(lo, "specs", {"left": 9.28, "top": 5.17, "width": 4.48, "height": 1.62})
        if img:
            _add_picture(slide, img, il["left"], il["top"], il["width"], il.get("height", 6.5), il.get("rotation", 0))
        _add_specs(slide, specs, sl["left"], sl["top"], sl["width"], sl.get("height", 1.62), font_pt=font_pt)

    _add_logo(slide, logo_path, pos=logo_pos)


def build_catalog_pptx(slides_data: list[dict], logo_path: str | None = None, logo_pos: dict | None = None, specs_font_pt: float | None = None) -> bytes:
    prs = Presentation()
    prs.slide_width  = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    font_pt = specs_font_pt if specs_font_pt else 10

    for sd in slides_data:
        _add_garment_slide(prs, sd, logo_path=logo_path, logo_pos=logo_pos, font_pt=font_pt)

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf.getvalue()
